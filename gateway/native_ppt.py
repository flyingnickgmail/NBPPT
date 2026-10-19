"""Bounded, passive OOXML validation, shared by rendering and settlement."""
from pathlib import Path, PurePosixPath
import stat
import zipfile
import xml.etree.ElementTree as ET

from pptx import Presentation


def walk_shapes(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:
            yield from walk_shapes(shape.shapes)


def validate_native_pptx(path: Path, requested_count=0):
    if path.is_symlink() or not path.is_file() or path.stat().st_nlink != 1:
        raise ValueError("PPTX 必须是任务内的独立普通文件")
    if not 1000 <= path.stat().st_size <= 64 * 1024 * 1024:
        raise ValueError("PPTX 文件大小不合法")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(x.file_size for x in entries) > 200 * 1024 * 1024:
            raise ValueError("PPTX 解压规模过大")
        names = [x.filename for x in entries]
        if len(set(names)) != len(names):
            raise ValueError("PPTX 包含重复文件项")
        for entry in entries:
            name = entry.filename.lower()
            parts = PurePosixPath(name)
            if parts.is_absolute() or ".." in parts.parts or "\\" in name or stat.S_ISLNK(entry.external_attr >> 16):
                raise ValueError("PPTX 包含不安全路径")
            if any(x in name for x in ("vbaproject", "activex", "embeddings/", "oleobject")):
                raise ValueError("PPTX 不允许宏或嵌入执行对象")
            if entry.flag_bits & 1:
                raise ValueError("PPTX 不允许加密内容")
            if name.endswith((".rels", ".xml", ".svg")):
                if entry.file_size > 4 * 1024 * 1024:
                    raise ValueError("PPTX 单个 XML 资源过大")
                raw = archive.read(entry)
                # XML can use UTF-16/32; ASCII-only byte matching misses its DTD.
                declarations = raw.replace(b'\x00', b'').upper()
                if b"<!DOCTYPE" in declarations or b"<!ENTITY" in declarations:
                    raise ValueError("PPTX 不允许文档实体")
                tree = ET.fromstring(raw)
                for node in tree.iter():
                    tag = node.tag.rsplit("}", 1)[-1].lower()
                    values = (node.get("ContentType", "") + node.get("Type", "")).lower()
                    if tag in ("oleobj", "control", "script", "foreignobject") or any(x in values for x in ("macroenabled", "vbaproject", "activex", "oleobject")):
                        raise ValueError("PPTX 不允许活动内容")
                    if name.endswith('.svg'):
                        if tag == 'style':
                            raise ValueError("PPTX 内嵌 SVG 不允许样式脚本加载外部资源")
                        for attr, value in node.attrib.items():
                            key = attr.rsplit('}', 1)[-1].lower()
                            if key == 'base' or key.startswith('on') or (key in ('href', 'src') and not value.startswith('#')) or 'url(' in value.lower():
                                raise ValueError("PPTX 内嵌 SVG 不允许脚本或外部资源")
            if name.endswith(".rels"):
                if any(x.get("TargetMode", "").lower() == "external" for x in tree):
                    raise ValueError("PPTX 不允许外部资源；引用网址可使用普通文字")
        if archive.testzip() is not None:
            raise ValueError("PPTX 压缩包校验失败")
    deck = Presentation(path)
    count = len(deck.slides)
    if not 1 <= count <= 60 or (requested_count and count != requested_count):
        raise ValueError("PPTX 页数与请求不一致")
    if abs(deck.slide_width / deck.slide_height - 16 / 9) > .01:
        raise ValueError("PPTX 必须为16:9宽屏")
    for page in deck.slides:
        if not any((shape.has_text_frame and shape.text.strip()) or shape.shape_type == 13 or shape.has_chart or shape.has_table for shape in walk_shapes(page.shapes)):
            raise ValueError("PPTX 包含空白页")
    return deck
