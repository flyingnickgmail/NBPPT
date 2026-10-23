"""Deterministic delivery gate shared by the worker and settlement boundary.

This checks files, not the truth of slide content. Published output is immutable.
"""
import hashlib
import json
import pathlib
import re
import subprocess
import zipfile
import base64
from fractions import Fraction
import html as html_lib
from html.parser import HTMLParser

from PIL import Image
from gateway.video_contract import voice_provider
from pptx import Presentation


class DeliveryError(ValueError):
    pass


def safe_job_id(job_id: str) -> str:
    if not re.fullmatch(r"(?:ppt|video|composite)_[a-f0-9]{10}", job_id):
        raise DeliveryError("Invalid job identifier")
    return job_id


def contained_file(root: pathlib.Path, relative: str) -> pathlib.Path:
    root = root.resolve()
    path = root / relative
    if path.is_symlink() or not path.resolve().is_relative_to(root) or not path.is_file():
        raise DeliveryError("Missing or unsafe artifact")
    if path.stat().st_size == 0:
        raise DeliveryError("Empty artifact")
    return path


class ImageReferences(HTMLParser):
    def __init__(self):
        super().__init__()
        self.sources = []

    def handle_starttag(self, tag, attrs):
        if tag == "img":
            self.sources.append(dict(attrs).get("src", ""))


def inspect_delivery(job_id: str, task_type: str, spec: dict, root: pathlib.Path) -> dict:
    root = root.resolve()
    safe_job_id(job_id)
    if task_type not in ("ppt", "video", "composite"):
        raise DeliveryError("Unsupported task type")
    files = []
    count = 0
    if task_type in ("ppt", "composite"):
        count = len(spec.get("slide_outline", []))
        if count < 1 or spec.get("slide_count") != count:
            raise DeliveryError("Invalid planned slide count")
        requested = spec.get("requested_slide_count", 0)
        if requested > 0 and requested != count:
            raise DeliveryError("Fixed slide count was not satisfied")
        deck = contained_file(root, f"{job_id}_deck.pptx")
        if spec.get("presentation_format") == "native-pptx-v1":
            from gateway.native_ppt import validate_native_pptx
            validate_native_pptx(deck, count)
        try:
            with zipfile.ZipFile(deck) as archive:
                if archive.testzip() is not None:
                    raise DeliveryError("PPTX archive integrity failure")
            prs = Presentation(deck)
            if len(prs.slides) != count:
                raise DeliveryError("PPTX page count mismatch")
            if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.01:
                raise DeliveryError("PPTX must be widescreen")
            if spec.get("generator_engine") == "claude_codex" and any(
                    not page.has_notes_slide or page.notes_slide.notes_text_frame.text != item.get("speaker_notes_cn")
                    for page, item in zip(prs.slides, spec["slide_outline"])):
                raise DeliveryError("Claude 讲稿没有完整写入 PPT 备注")
        except Exception as exc:
            raise DeliveryError("Invalid PPTX or page count/aspect ratio") from exc
        expected = [f"assets/slide_preview_{i}.png" for i in range(1, count + 1)]
        found = {p.relative_to(root).as_posix() for p in (root / "assets").glob("slide_preview_*.png")}
        if found != set(expected):
            raise DeliveryError("Preview page count or numbering mismatch")
        for name in expected:
            path = contained_file(root, name)
            try:
                with Image.open(path) as im:
                    im.load()
                    if im.format != "PNG" or im.size != (1672, 941):
                        raise DeliveryError("Preview must be a 1672x941 PNG")
            except Exception as exc:
                raise DeliveryError("Invalid preview image or resolution") from exc
            files.append(path)
        html = contained_file(root, f"{job_id}_deck_interactive.html")
        refs = ImageReferences()
        refs.feed(html.read_text(encoding="utf-8"))
        embedded = ["data:image/png;base64," + base64.b64encode((root / name).read_bytes()).decode("ascii") for name in expected]
        if refs.sources != expected and refs.sources != embedded:
            raise DeliveryError("HTML must reference every preview in page order")
        files.extend([deck, html])
        if spec.get("generator_engine") == "claude_codex":
            from gateway.content_plan import read_plan, validate_content_plan, validate_content_handoff, outline_markdown
            plan_path = contained_file(root, "content_plan.json")
            text_path = contained_file(root, "content_outline.md")
            plan = validate_content_plan(read_plan(plan_path), spec)
            if plan != spec.get("claude_content_plan"):
                raise DeliveryError("交付策划与实际 Claude 策划不一致")
            validate_content_handoff(spec["slide_outline"], plan)
            if text_path.stat().st_size > 4 * 1024 * 1024 or text_path.read_text(encoding="utf-8") != outline_markdown(plan):
                raise DeliveryError("文字版缺失或与策划不一致")
            files.extend([plan_path, text_path])
    if task_type in ("video", "composite"):
        video = contained_file(root, f"{job_id}_video.mp4")
        try:
            result = subprocess.run(
                ["ffprobe", "-v", "error", "-show_streams", "-show_format", "-of", "json", str(video)],
                capture_output=True, timeout=30, check=True,
            )
            info = json.loads(result.stdout)
            duration = float(info["format"]["duration"])
            streams = [s for s in info["streams"] if s.get("codec_type") == "video" and s.get("width", 0) > 0 and s.get("height", 0) > 0]
            if duration <= 0 or not streams:
                raise ValueError("No video stream")
            expected_duration = spec.get("video_duration_sec")
            if expected_duration and abs(duration - float(expected_duration)) > 0.5:
                raise ValueError("Video duration mismatch")
            ratio = spec.get("video_aspect_ratio")
            if ratio in ("16:9", "9:16"):
                width, height = map(int, ratio.split(":"))
                if abs(streams[0]["width"] / streams[0]["height"] - width / height) > 0.02:
                    raise ValueError("Video aspect ratio mismatch")
            if spec.get('generation_contract') in ('skill-video-v1','story-video-v2'):
                expected_size = (1280, 720) if ratio == '16:9' else (720, 1280)
                stream = streams[0]
                if (len(streams) != 1 or (stream['width'], stream['height']) != expected_size
                    or stream.get('codec_name') != 'h264' or stream.get('pix_fmt') != 'yuv420p'
                    or Fraction(stream.get('avg_frame_rate', '0')) != 24
                    or abs(duration - expected_duration) > 1 / 24 + .001
                    or (spec.get('video_width'), spec.get('video_height')) != expected_size
                    or spec.get('fps') != 24):
                    raise ValueError('Video encoding contract mismatch')
                audio = [s for s in info['streams'] if s.get('codec_type') == 'audio']
                if type(spec.get('video_audio')) is not bool or bool(audio) != spec['video_audio']:
                    raise ValueError('Audio metadata mismatch')
                narration = spec.get('narration', {})
                if spec.get('video_narration') is True and (not audio or narration.get('status') != 'generated'):
                    raise ValueError('Requested narration was not delivered')
                if spec.get('video_narration') is True and 'video_voice' in spec:
                    voice = spec['video_voice']
                    if (narration.get('voice') != voice or
                            narration.get('provider') != voice_provider(voice)):
                        raise ValueError('Requested voice was not delivered')
                if spec.get('video_subtitles') is True and spec.get('subtitles_burned') is not True:
                    raise ValueError('Requested burned subtitles were not delivered')
                if audio and (len(audio) != 1 or audio[0].get('codec_name') != 'aac'
                              or abs(float(audio[0].get('duration', 0)) - expected_duration) > .05):
                    raise ValueError('Audio encoding or duration mismatch')
                if audio and spec.get('generation_contract')=='story-video-v2':
                    from gateway.audio_contract import verify_loudness
                    verify_loudness(video)
                # Decode every frame: a valid header alone does not prove a playable file.
                subprocess.run(['ffmpeg', '-v', 'error', '-xerror', '-nostdin', '-threads', '2',
                                '-i', str(video), '-f', 'null', '-'],
                               capture_output=True, timeout=120, check=True)
                poster = contained_file(root, 'assets/video_poster.png')
                with Image.open(poster) as im:
                    im.load()
                    if im.format != 'PNG' or im.size != expected_size:
                        raise ValueError('Video poster size mismatch')
                captions = contained_file(root, 'subtitles.vtt')
                scenes = spec.get('video_scenes')
                if not isinstance(scenes, list) or not 1 <= len(scenes) <= (30 if spec['video_duration_sec']>60 else 10):
                    raise ValueError('Missing storyboard')
                cursor = 0
                expected_cues = []
                def stamp(seconds):
                    millis = round(seconds * 1000)
                    return f'{millis // 3600000:02}:{millis % 3600000 // 60000:02}:{millis % 60000 // 1000:02}.{millis % 1000:03}'
                speech_scenes = narration.get('scenes', [])
                if audio and len(speech_scenes) != len(scenes):
                    raise ValueError('Missing per-scene narration evidence')
                for index, scene in enumerate(scenes):
                    length = scene['duration_sec']
                    if type(length) is not int or length < 3:
                        raise ValueError('Invalid scene duration')
                    end = cursor + length
                    cue_end = end
                    if audio:
                        speech = speech_scenes[index]
                        spoken = speech.get('speech_duration_sec', 0)
                        if speech.get('scene_num') != index + 1 or not 0 < spoken <= length or not 1 <= speech.get('speed', 0) <= 1.2:
                            raise ValueError('Narration timing exceeds scene')
                        cue_end = cursor + speech.get("caption_duration_sec",spoken)
                        if cue_end not in (end,cursor+spoken):
                            raise ValueError("Invalid caption duration")
                    expected_cues.append(f'{stamp(cursor)} --> {stamp(cue_end)}')
                    cursor = end
                text = captions.read_text(encoding='utf-8')
                cues = re.findall(r'^\d{2}:\d{2}:\d{2}\.\d{3} --> \d{2}:\d{2}:\d{2}\.\d{3}$', text, re.M)
                if not text.startswith('WEBVTT\n') or cues != expected_cues or cursor != expected_duration:
                    raise ValueError('Subtitles must cover every scene exactly')
                # Each cue must contain its corresponding visible content, with no empty scene.
                blocks = text.strip().split('\n\n')[1:]
                if len(blocks) != len(scenes):
                    raise ValueError('Subtitle scene count mismatch')
                for block, scene in zip(blocks, scenes):
                    expected_text = scene.get('narration') if audio or spec.get('subtitle_content') == 'narration' else scene['title'] + '\n' + scene['body']
                    if not expected_text or expected_text not in html_lib.unescape(block):
                        raise ValueError('Subtitle content differs from storyboard')
                files.extend([poster, captions])
        except Exception as exc:
            raise DeliveryError("Video cannot be validated; ffprobe and valid media are required") from exc
        if spec.get('video_visual_mode')=='live_action':
            from gateway.live_action_contract import validate_sources
            from gateway.footage_contract import credits_html
            sources_path=contained_file(root,'footage_sources.json')
            if sources_path.stat().st_size>512*1024: raise DeliveryError('Footage provenance too large')
            clips=validate_sources(json.loads(sources_path.read_text()),spec)
            credits=contained_file(root,'footage_credits.html')
            if credits.read_text()!=credits_html(clips): raise DeliveryError('Footage credits mismatch')
            files.extend([sources_path,credits])
        if spec.get('planner_provider') == 'claude_cli':
            from gateway.video_manuscript import validate_manuscript, validate_handoff, manuscript_markdown, read_json_text
            manuscript_path = contained_file(root,'manuscript.json')
            if manuscript_path.stat().st_size > 128*1024:
                raise DeliveryError('内容稿过大')
            manuscript = validate_manuscript(read_json_text(manuscript_path.read_text()),spec)
            if manuscript != spec.get('video_manuscript'):
                raise DeliveryError('文字稿与实际规划不一致')
            story_path = contained_file(root,'storyboard.json')
            validate_handoff(manuscript, read_json_text(story_path.read_text()),spec)
            expected_plan = {'task_summary_cn':spec['task_summary_cn'],'story':spec['story'],'scenes':spec['video_scenes']}
            if read_json_text(story_path.read_text()) != expected_plan:
                raise DeliveryError('分镜稿与成片计划不一致')
            text_path = contained_file(root,'content_manuscript.md')
            voice_path = contained_file(root,'narration.txt')
            if (text_path.read_text()!=manuscript_markdown(manuscript) or
                    voice_path.read_text()!='\n\n'.join(s['narration'] for s in manuscript['sections'])):
                raise DeliveryError('下载文字与实际口播不一致')
            sheet=contained_file(root,'assets/storyboard.png')
            with Image.open(sheet) as image:
                image.verify()
            files.extend([manuscript_path,story_path,text_path,voice_path,sheet])
        files.append(video)
    return {
        "version": 1, "job_id": job_id, "task_type": task_type,
        "planned_slides": count, "file_status": "PASSED", "content_status": "NOT_REVIEWED",
        "files": [
            {"path": p.relative_to(root).as_posix(), "size": p.stat().st_size,
             "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
            for p in sorted(files)
        ],
    }


def verify_published(job_id: str, task_type: str, spec: dict, root: pathlib.Path) -> dict:
    manifest = json.loads(contained_file(root, "delivery.json").read_text())
    actual = inspect_delivery(job_id, task_type, spec, root)
    if manifest != actual:
        raise DeliveryError("Published delivery differs from verified manifest")
    return actual


def delivery_urls(job_id: str, manifest: dict) -> dict:
    base = f"/artifacts/{job_id}/"
    paths = [f["path"] for f in manifest["files"]]
    return {
        "pptx_url": next((base + p for p in paths if p.endswith(".pptx")), None),
        "html_deck_url": next((base + p for p in paths if p.endswith(".html") and p != "footage_credits.html"), None),
        "content_outline_url": base + 'content_outline.md' if 'content_outline.md' in paths else None,
        "content_plan_url": base + 'content_plan.json' if 'content_plan.json' in paths else None,
        "footage_sources_url": base+'footage_sources.json' if 'footage_sources.json' in paths else None,
        "footage_credits_url": base+'footage_credits.html' if 'footage_credits.html' in paths else None,
        "manuscript_url": base+'content_manuscript.md' if 'content_manuscript.md' in paths else None,
        "narration_url": base+'narration.txt' if 'narration.txt' in paths else None,
        "storyboard_url": base+'storyboard.json' if 'storyboard.json' in paths else None,
        "storyboard_preview_url": base+'assets/storyboard.png' if 'assets/storyboard.png' in paths else None,
        "mp4_url": next((base + p for p in paths if p.endswith(".mp4")), None),
        "video_poster_url": base + 'assets/video_poster.png' if 'assets/video_poster.png' in paths else None,
        "subtitles_url": base + 'subtitles.vtt' if 'subtitles.vtt' in paths else None,
        "slide_preview_urls": [base + f"assets/slide_preview_{i}.png" for i in range(1, manifest["planned_slides"] + 1)],
        "chart_urls": [], "keyframe_urls": [],
    }
