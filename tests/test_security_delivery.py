"""Offline regression tests: no server access, CLI generation or model billing."""
import asyncio
import hashlib
import json
import pathlib
import sqlite3
import subprocess
import shutil

import httpx
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from gateway.billing_tcc import BillingManager
from gateway.delivery import DeliveryError, inspect_delivery, verify_published, delivery_urls
from gateway import main
from gateway import mcp_server
from worker import worker_api


JOB = "ppt_0123456789"
KEY_A = "test-a-" + "a" * 40
KEY_B = "test-b-" + "b" * 40
ADMIN = "test-admin-" + "c" * 40


def run(coro):
    return asyncio.run(coro)


@pytest.fixture
def manager(tmp_path, monkeypatch):
    monkeypatch.delenv("NBPPT_ADMIN_API_KEY", raising=False)
    monkeypatch.delenv("NBPPT_PUBLIC_ORIGIN", raising=False)
    monkeypatch.setenv("NBPPT_COOKIE_SECURE", "true")
    mgr = BillingManager(str(tmp_path / "test.db"), tmp_path / "jobs")
    run(mgr.initialize())
    with sqlite3.connect(mgr.db_path) as db:
        for user, key, tier in (("a", KEY_A, "studio"), ("b", KEY_B, "studio"), ("admin", ADMIN, "admin")):
            db.execute("INSERT INTO users(id,phone_or_email,tier,credits_available,created_at) VALUES (?,?,?,?,?)", (user,user,tier,2000,"now"))
            db.execute("INSERT INTO mcp_api_keys(api_key,user_id,created_at,scope) VALUES (?,?,?,?)", (key,user,"now","admin" if tier=="admin" else "generate"))
    monkeypatch.setattr(main, "billing", mgr)
    monkeypatch.setattr(main, "JOBS_ROOT", mgr.jobs_root)
    return mgr


def make_delivery(root, job=JOB, count=2, publish=True):
    root.mkdir(parents=True, exist_ok=True)
    (root / "assets").mkdir(exist_ok=True)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(16), Inches(9)
    for i in range(1, count+1):
        Image.new("RGB", (1672,941), (i*40,60,120)).save(root / f"assets/slide_preview_{i}.png")
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(str(root/f"assets/slide_preview_{i}.png"),0,0,width=deck.slide_width,height=deck.slide_height)
    deck.save(root / f"{job}_deck.pptx")
    (root/f"{job}_deck_interactive.html").write_text("<html>"+"".join(f'<img src="assets/slide_preview_{i}.png">' for i in range(1,count+1))+"</html>")
    spec = {"slide_count":count,"requested_slide_count":count,"slide_outline":[{} for _ in range(count)]}
    manifest = inspect_delivery(job,"ppt",spec,root)
    if publish:
        (root/"delivery.json").write_text(json.dumps(manifest))
    return spec, manifest


async def prepare(manager, job=JOB):
    await manager.try_freeze_credits(KEY_A, job, "ppt", "test", 60)
    spec, manifest = make_delivery(manager.jobs_root / job / "output", job)
    await manager.save_compiled_spec(job,spec)
    return manifest


def balance(manager):
    with sqlite3.connect(manager.db_path) as db:
        return db.execute("SELECT credits_available,credits_frozen FROM users WHERE id='a'").fetchone()


@pytest.mark.parametrize("defect", ["missing", "empty", "corrupt_pptx", "missing_preview", "extra_preview", "small_preview", "corrupt_preview", "broken_html", "reordered_html", "fixed_count", "pptx_count", "changed_manifest", "symlink"])
def test_strict_gate_rejects_bad_deliveries(tmp_path, defect):
    root = tmp_path / "output"
    spec, _ = make_delivery(root)
    pptx = root/f"{JOB}_deck.pptx"
    img = root/"assets/slide_preview_1.png"
    if defect == "missing": pptx.unlink()
    elif defect == "empty": pptx.write_bytes(b"")
    elif defect == "corrupt_pptx": pptx.write_bytes(b"not a zip")
    elif defect == "missing_preview": img.unlink()
    elif defect == "extra_preview": (root/"assets/slide_preview_3.png").write_bytes(img.read_bytes())
    elif defect == "small_preview": Image.new("RGB",(320,180)).save(img)
    elif defect == "corrupt_preview": img.write_bytes(b"broken")
    elif defect == "broken_html": (root/f"{JOB}_deck_interactive.html").write_text('<img src="missing.png">')
    elif defect == "reordered_html": (root/f"{JOB}_deck_interactive.html").write_text('<img src="assets/slide_preview_2.png"><img src="assets/slide_preview_1.png">')
    elif defect == "fixed_count": spec["requested_slide_count"] = 3
    elif defect == "pptx_count":
        deck=Presentation(pptx); deck.slides.add_slide(deck.slide_layouts[6]); deck.save(pptx)
    elif defect == "changed_manifest": (root/"delivery.json").write_text('{}')
    elif defect == "symlink":
        other=tmp_path/"other.png"; img.rename(other); img.symlink_to(other)
    with pytest.raises((DeliveryError, ValueError)):
        verify_published(JOB,"ppt",spec,root)


def test_valid_delivery_and_video_separation(tmp_path, monkeypatch):
    spec,manifest=make_delivery(tmp_path/"ppt")
    assert verify_published(JOB,"ppt",spec,tmp_path/"ppt") == manifest
    assert manifest["content_status"] == "NOT_REVIEWED"
    video=tmp_path/"video";video.mkdir();(video/"video_0123456789_video.mp4").write_bytes(b"video")
    class Probe:
        stdout=b'{"format":{"duration":"2"},"streams":[{"codec_type":"video","width":1920,"height":1080}]}'
    monkeypatch.setattr("gateway.delivery.subprocess.run",lambda *a,**k: Probe())
    assert inspect_delivery("video_0123456789","video",{},video)["planned_slides"] == 0
    with pytest.raises(DeliveryError): inspect_delivery("composite_0123456789","composite",spec,video)


def test_no_artifact_no_settlement(manager):
    async def scenario():
        await manager.try_freeze_credits(KEY_A,JOB,"ppt","test",60)
        await manager.save_compiled_spec(JOB,{"slide_count":2,"slide_outline":[{},{}]})
        with pytest.raises(DeliveryError): await manager.confirm_settle(JOB,.1,{})
        await manager.cancel_refund(JOB,"missing delivery")
    run(scenario())
    assert balance(manager)==(2000,0)
    assert run(manager.get_job(JOB,"a"))["status"]=="FAILED"


def test_repeated_settlement_and_refund(manager):
    async def scenario():
        manifest=await prepare(manager)
        results=await asyncio.gather(*[manager.confirm_settle(JOB,.1,delivery_urls(JOB,manifest)) for _ in range(5)])
        assert sum(bool(r.get("already_settled")) for r in results)==4
        assert (await manager.cancel_refund(JOB,"late error"))["refunded"]==0
    run(scenario())
    assert balance(manager)==(1958,0)
    with sqlite3.connect(manager.db_path) as db:
        assert db.execute("SELECT count(*) FROM credit_ledger WHERE action='SETTLE'").fetchone()[0]==1


def test_settlement_refund_race(manager):
    async def scenario():
        manifest=await prepare(manager)
        await asyncio.gather(manager.confirm_settle(JOB,.1,delivery_urls(JOB,manifest)),manager.cancel_refund(JOB,"cancel"),return_exceptions=True)
    run(scenario())
    job=run(manager.get_job(JOB,"a"))
    assert balance(manager)==((1958,0) if job["status"]=="COMPLETED" else (2000,0))
    with sqlite3.connect(manager.db_path) as db:
        assert db.execute("SELECT count(*) FROM credit_ledger WHERE action IN ('SETTLE','REFUND')").fetchone()[0]==1


def test_restart_refunds_once(manager):
    async def scenario():
        await manager.try_freeze_credits(KEY_A,JOB,"ppt","test",60)
        await manager.recover_interrupted();await manager.recover_interrupted()
    run(scenario());assert balance(manager)==(2000,0)


def test_public_and_user_boundaries(manager, monkeypatch):
    async def scenario():
        manifest=await prepare(manager)
        await manager.confirm_settle(JOB,.1,delivery_urls(JOB,manifest))
        transport=httpx.ASGITransport(app=main.app)
        async with httpx.AsyncClient(transport=transport,base_url="https://testserver") as c:
            assert (await c.get('/health')).json()=={"status":"alive"}
            for path in ('/api/jobs','/api/wallet','/api/settings/models',f'/artifacts/{JOB}/{JOB}_deck.pptx'):
                assert (await c.get(path)).status_code==401
            for path in ('/api/jobs/stream','/api/wallet/topup','/api/settings/models','/mcp'):
                assert (await c.post(path,json={})).status_code==401
            assert (await c.get('/api/wallet?api_key='+KEY_A)).status_code==401
            assert (await c.get('/api/jobs',headers={'Authorization':'Bearer sk-mcp-demo-pro-8899'})).status_code==401
            h={'Authorization':'Bearer '+KEY_A}
            assert (await c.get('/api/jobs',headers=h)).json()['jobs'][0]['job_id']==JOB
            assert (await c.get('/api/jobs',headers={'Authorization':'Bearer '+KEY_B})).json()['jobs']==[]
            assert (await c.get(f'/artifacts/{JOB}/{JOB}_deck.pptx',headers={'Authorization':'Bearer '+KEY_B})).status_code==404
            assert (await c.get(f'/artifacts/{JOB}/{JOB}_deck.pptx',headers=h)).status_code==200
            assert (await c.get(f'/artifacts/{JOB}/delivery.json',headers=h)).status_code==404
            assert (await c.post('/api/settings/models',headers=h,json={})).status_code==403
            assert (await c.post('/api/wallet/topup',headers=h,json={'amount':500})).status_code==403
            assert KEY_A not in (await c.get('/api/wallet',headers=h)).text
            login=await c.post('/api/session',json={'api_key':KEY_A},headers={'Origin':'https://testserver'})
            assert login.status_code==200
            assert 'HttpOnly' in login.headers['set-cookie'] and 'Secure' in login.headers['set-cookie']
            assert (await c.get(f'/artifacts/{JOB}/assets/slide_preview_1.png')).status_code==200
            html=await c.get(f'/artifacts/{JOB}/{JOB}_deck_interactive.html')
            assert 'sandbox allow-scripts;' in html.headers['content-security-policy']
            assert 'allow-same-origin' not in html.headers['content-security-policy']
            assert (await c.post('/api/wallet/keys',json={},headers={'Origin':'https://evil.example'})).status_code==403
            assert (await c.delete('/api/session',headers={'Origin':'https://testserver'})).status_code==200
            assert (await c.get('/api/jobs')).status_code==401
    run(scenario())


def test_revoked_and_expired_credentials(manager):
    async def scenario():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app),base_url='https://testserver') as c:
            await c.post('/api/session',json={'api_key':KEY_A})
            with sqlite3.connect(manager.db_path) as db:db.execute('UPDATE web_sessions SET expires_at=0')
            assert (await c.get('/api/jobs')).status_code==401
            await c.post('/api/session',json={'api_key':KEY_A})
            with sqlite3.connect(manager.db_path) as db:db.execute('UPDATE mcp_api_keys SET is_active=0 WHERE api_key=?',(KEY_A,))
            assert (await c.get('/api/jobs')).status_code==401
            assert (await c.get('/api/jobs',headers={'Authorization':'Bearer '+KEY_A})).status_code==401
    run(scenario())


def test_worker_denies_missing_or_default_secret(tmp_path,monkeypatch):
    monkeypatch.setattr(worker_api,'WORKSPACE_JOBS_DIR',tmp_path)
    async def scenario():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=worker_api.app),base_url='http://testserver') as c:
            payload={'job_id':JOB,'english_spec':{}}
            for secret in ('','super-secret-mtls-token-2026'):
                monkeypatch.setattr(worker_api,'WORKER_SECRET',secret)
                assert (await c.post('/internal/execute-stream',json=payload,headers={'X-Worker-Secret':secret})).status_code==403
            monkeypatch.setattr(worker_api,'WORKER_SECRET','x'*40)
            payload['job_id']='../escape'
            assert (await c.post('/internal/execute-stream',json=payload,headers={'X-Worker-Secret':'x'*40})).status_code==422
    run(scenario());assert list(tmp_path.iterdir())==[]


def test_global_images_are_not_collected(tmp_path, monkeypatch):
    global_dir=tmp_path/'global';global_dir.mkdir()
    Image.new('RGB',(1672,941)).save(global_dir/'other.png')
    monkeypatch.setattr(worker_api,'CODEX_GENERATED_IMAGES_DIR',global_dir)
    assert worker_api._collect_new_codex_images(0,tmp_path/'assets',2)==[]


def test_gateway_rejects_fake_worker_completion(manager,monkeypatch):
    real_client=httpx.AsyncClient
    class Stream:
        async def __aenter__(self):return self
        async def __aexit__(self,*args):pass
        def raise_for_status(self):pass
        async def aiter_lines(self):yield 'data: '+json.dumps({'type':'completed','total_cost_usd':.1,'cdn_urls':{}})
    class WorkerClient:
        def __init__(self,*args,**kwargs):pass
        async def __aenter__(self):return self
        async def __aexit__(self,*args):pass
        def stream(self,*args,**kwargs):return Stream()
    async def compile_spec(*args,**kwargs):return {'slide_count':2,'slide_outline':[{},{}]}
    monkeypatch.setattr(main,'compile_cn_request_to_english_spec',compile_spec)
    monkeypatch.setattr(mcp_server,'compile_cn_request_to_english_spec',compile_spec)
    monkeypatch.setattr(main.httpx,'AsyncClient',WorkerClient)
    async def scenario():
        async with real_client(transport=httpx.ASGITransport(app=main.app),base_url='https://testserver') as c:
            h={'Authorization':'Bearer '+KEY_A}
            r=await c.post('/api/jobs/stream',headers=h,json={'prompt_cn':'test','slide_count':2})
            events=[json.loads(line[6:]) for line in r.text.splitlines() if line.startswith('data: ')]
            assert events[-1]['type']=='error'
            assert not any(e['type']=='completed' for e in events)
            jobs=mcp_server.PresentationJobs(manager)
            result=await jobs.create({'api_key':KEY_A,'user_id':'a'},'test','auto',2,'fake-worker-test')
            await asyncio.gather(*list(jobs.tasks))
            assert (await jobs.status({'user_id':'a'},result['job_id']))['status']=='FAILED'

    run(scenario());assert balance(manager)==(2000,0)


def test_worker_success_publishes_only_verified_files(tmp_path, monkeypatch):
    monkeypatch.setattr(worker_api,'WORKSPACE_JOBS_DIR',tmp_path)
    monkeypatch.setattr(worker_api,'WORKER_SECRET','s'*40)
    monkeypatch.setattr(worker_api,'CODEX_BIN',__file__)
    class Process:
        returncode=0
    async def fake_exec(*args,**kwargs):
        # A deterministic local stand-in for the paid image generator.
        assets=pathlib.Path(kwargs['cwd'])/'staging/assets'
        for i in (1,2): Image.new('RGB',(1672,941),(i*40,60,120)).save(assets/f'slide_preview_{i}.png')
        assert 'NBPPT_ADMIN_API_KEY' not in kwargs['env']
        assert 'OVERSEAS_WORKER_SECRET' not in kwargs['env']
        return Process()
    monkeypatch.setattr(worker_api.asyncio,'create_subprocess_exec',fake_exec)
    spec={'slide_count':2,'slide_outline':[{},{}],'requested_slide_count':2}
    async def scenario():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=worker_api.app),base_url='http://testserver') as c:
            payload={'job_id':JOB,'english_spec':spec,'runtime_keys':{'account_safe_mode':False}}
            r=await c.post('/internal/execute-stream',json=payload,headers={'X-Worker-Secret':'s'*40})
            events=[json.loads(line[6:]) for line in r.text.splitlines() if line.startswith('data: ')]
            assert events[-1]['type']=='completed',events
            assert events[-1]['qa_report']['content_status']=='NOT_REVIEWED'
            assert not (tmp_path/JOB/'staging').exists()
            assert verify_published(JOB,'ppt',spec,tmp_path/JOB/'output')
            assert (await c.post('/internal/execute-stream',json=payload,headers={'X-Worker-Secret':'s'*40})).status_code==409
    run(scenario())


def test_worker_failed_cli_never_completes(tmp_path,monkeypatch):
    monkeypatch.setattr(worker_api,'WORKSPACE_JOBS_DIR',tmp_path)
    monkeypatch.setattr(worker_api,'WORKER_SECRET','s'*40)
    monkeypatch.setattr(worker_api,'CODEX_BIN',__file__)
    class Process:
        returncode=1
    async def fake_exec(*args,**kwargs):return Process()
    monkeypatch.setattr(worker_api.asyncio,'create_subprocess_exec',fake_exec)
    async def scenario():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=worker_api.app),base_url='http://testserver') as c:
            r=await c.post('/internal/execute-stream',headers={'X-Worker-Secret':'s'*40},json={'job_id':JOB,'english_spec':{'slide_count':2,'slide_outline':[{},{}]},'runtime_keys':{'account_safe_mode':False}})
            assert '"type": "completed"' not in r.text
            assert '"type": "error"' in r.text
    run(scenario());assert not (tmp_path/JOB/'output').exists()


def test_atomic_freeze_prevents_overspending(manager):
    with sqlite3.connect(manager.db_path) as db:db.execute("UPDATE users SET credits_available=60 WHERE id='a'")
    async def scenario():
        results=await asyncio.gather(manager.try_freeze_credits(KEY_A,JOB,'ppt','test',60),manager.try_freeze_credits(KEY_A,'ppt_1111111111','ppt','test',60),return_exceptions=True)
        assert sum(isinstance(r,ValueError) for r in results)==1
    run(scenario());assert balance(manager)==(0,60)


def test_manifest_mismatch_prevents_settlement(manager):
    async def scenario():
        manifest=await prepare(manager)
        with pytest.raises(ValueError):await manager.confirm_settle(JOB,.1,{'pptx_url':'https://evil.example/fake.pptx'})
        (manager.jobs_root/JOB/'output'/f'{JOB}_deck_interactive.html').write_text('tampered')
        with pytest.raises(DeliveryError):await manager.confirm_settle(JOB,.1,delivery_urls(JOB,manifest))
    run(scenario());assert balance(manager)==(1940,60)


def test_settled_artifact_tampering_is_not_served(manager):
    async def scenario():
        manifest=await prepare(manager)
        await manager.confirm_settle(JOB,.1,delivery_urls(JOB,manifest))
        (manager.jobs_root/JOB/'output'/f'{JOB}_deck.pptx').write_bytes(b'changed')
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app),base_url='https://testserver') as c:
            r=await c.get(f'/artifacts/{JOB}/{JOB}_deck.pptx',headers={'Authorization':'Bearer '+KEY_A})
            assert r.status_code==404
    run(scenario())


def test_admin_settings_and_topup(manager,monkeypatch):
    saved=[]
    monkeypatch.setattr(main,'save_runtime_settings',lambda values:saved.append(values))
    async def settings():return {'configured':True}
    monkeypatch.setattr(main,'get_model_settings',settings)
    async def scenario():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app),base_url='https://testserver') as c:
            h={'Authorization':'Bearer '+ADMIN}
            r=await c.post('/api/settings/models',headers=h,json={'claude_model':'test-model'})
            assert r.status_code==200
            assert saved==[{'claude_model':'test-model'}]
            r=await c.post('/api/wallet/topup',headers=h,json={'amount':10})
            assert r.status_code==200
            assert r.json()['wallet']['user']['credits_available']==2010
    run(scenario())


def test_browser_disconnect_refunds(manager,monkeypatch):
    async def cancelled_compile(*args,**kwargs):raise asyncio.CancelledError()
    monkeypatch.setattr(main,'compile_cn_request_to_english_spec',cancelled_compile)
    class State:principal={'api_key':KEY_A}
    class Request:state=State()
    async def scenario():
        response=await main.stream_create_job(main.CreateJobRequest(prompt_cn='test'),Request())
        with pytest.raises(asyncio.CancelledError):
            async for _ in response.body_iterator:pass
    run(scenario());assert balance(manager)==(2000,0)


def test_real_video_validation(tmp_path):
    if not shutil.which('ffmpeg') or not shutil.which('ffprobe'):
        pytest.skip('ffmpeg/ffprobe not installed')
    job='video_0123456789';video=tmp_path/f'{job}_video.mp4'
    subprocess.run(['ffmpeg','-v','error','-f','lavfi','-i','color=c=blue:s=192x108:r=10','-t','1','-c:v','mpeg4',str(video)],check=True,timeout=30)
    assert inspect_delivery(job,'video',{'video_duration_sec':1,'video_aspect_ratio':'16:9'},tmp_path)['file_status']=='PASSED'
    with pytest.raises(DeliveryError):inspect_delivery(job,'video',{'video_duration_sec':9},tmp_path)
    video.write_bytes(b'not a video')
    with pytest.raises(DeliveryError):inspect_delivery(job,'video',{},tmp_path)


def test_same_database_cannot_start_two_gateways(manager):
    async def scenario():
        async with main.lifespan(main.app):
            with pytest.raises(BlockingIOError):
                async with main.lifespan(main.app):pass
    run(scenario())


def test_empty_install_does_not_create_demo_user(tmp_path,monkeypatch):
    monkeypatch.delenv('NBPPT_ADMIN_API_KEY',raising=False)
    mgr=BillingManager(str(tmp_path/'fresh.db'))
    run(mgr.initialize())
    with sqlite3.connect(mgr.db_path) as db:
        assert db.execute('SELECT count(*) FROM mcp_api_keys').fetchone()[0]==0


def test_bootstrap_migration_preserves_balance_and_disables_demo(manager,monkeypatch):
    from gateway.config import DEFAULT_DEMO_API_KEY
    with sqlite3.connect(manager.db_path) as db:
        db.execute("UPDATE users SET phone_or_email='founder@personal-agent.cn' WHERE id='a'")
        db.execute("INSERT INTO mcp_api_keys(api_key,user_id,created_at) VALUES (?,'a','now')",(DEFAULT_DEMO_API_KEY,))
    monkeypatch.setenv('NBPPT_ADMIN_API_KEY','new-bootstrap-'+('x'*40))
    mgr=BillingManager(manager.db_path)
    run(mgr.initialize());assert balance(manager)==(2000,0)
    with sqlite3.connect(manager.db_path) as db:
        assert db.execute('SELECT is_active FROM mcp_api_keys WHERE api_key=?',(DEFAULT_DEMO_API_KEY,)).fetchone()[0]==0


def test_manifest_normalizes_symlinked_root(tmp_path):
    actual=tmp_path/'real';spec,manifest=make_delivery(actual)
    alias=tmp_path/'alias';alias.symlink_to(actual,target_is_directory=True)
    assert verify_published(JOB,'ppt',spec,alias)==manifest


def test_worker_cancellation_cleans_process_group(tmp_path,monkeypatch):
    monkeypatch.setattr(worker_api,'WORKSPACE_JOBS_DIR',tmp_path)
    monkeypatch.setattr(worker_api,'WORKER_SECRET','s'*40)
    monkeypatch.setattr(worker_api,'CODEX_BIN',__file__)
    killed=[]
    class Process:
        returncode=None
        pid=99999999
        async def wait(self):self.returncode=-9
    proc=Process()
    async def fake_exec(*args,**kwargs):return proc
    async def interrupted_sleep(*args):raise asyncio.CancelledError()
    monkeypatch.setattr(worker_api.asyncio,'create_subprocess_exec',fake_exec)
    monkeypatch.setattr(worker_api.asyncio,'sleep',interrupted_sleep)
    monkeypatch.setattr(worker_api.os,'killpg',lambda pid,sig:killed.append(pid))
    async def scenario():
        req=worker_api.JobExecutionRequest(job_id=JOB,english_spec={'slide_count':2,'slide_outline':[{},{}]},runtime_keys={'account_safe_mode':False})
        response=await worker_api.execute_stream(req,'s'*40)
        with pytest.raises(asyncio.CancelledError):
            async for _ in response.body_iterator:pass
    run(scenario())
    assert killed==[proc.pid] and proc.returncode==-9
    assert not (tmp_path/JOB/'output').exists()
