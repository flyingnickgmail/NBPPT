"""Hybrid admission, real process handoff failures, immutable content and text delivery."""
import asyncio
import copy
import json
import sqlite3
import sys
from types import SimpleNamespace

import httpx
import pytest
from pptx import Presentation

from gateway import main
from gateway.content_plan import validate_content_plan, validate_content_handoff, outline_markdown
from gateway.deepseek_compiler import compile_cn_request_to_english_spec, validate_direct_completion
from gateway.delivery import verify_published, delivery_urls
from gateway.mcp_server import PresentationJobs
from worker import claude_generation as claude, direct_generation as direct
from test_claude_generation import enabled
from test_direct_generation import spec, plan, fixture_output
from test_security_delivery import manager, KEY_A, KEY_B, balance


def content(count=5):
    value = {**plan(count), 'audience_cn':'经营决策者', 'central_thesis_cn':'先把责任闭环做实',
             'narrative_arc_cn':'问题到机制，到验证再到行动', 'evidence_boundary_cn':'分析判断需要试点验证'}
    for page in value['slide_outline']:
        page.update(argument_cn='机制推演与替代方案取舍', speaker_notes_cn='讲稿展开本页机制及其适用边界。', transition_cn='由此引出下一项行动。')
    return value


def hybrid_spec(count=5):
    return {**spec(count), 'generator_engine':'claude_codex', 'account_owner_user_id':'a',
            'generation_contract':'skill-direct-v1', 'claude_content_plan':content(count)}


def test_hybrid_contract_and_prompt_keep_argument_and_brand(enabled):
    value = asyncio.run(compile_cn_request_to_english_spec(spec()['original_prompt_cn'], 'ppt', {'execution_mode':'claude_codex'}))
    assert value['generator_engine']=='claude_codex' and value['slide_count']==0 and not value['slide_outline']
    assert value['matched_template']['id']=='auto' and 'presentation_format' not in value
    prompt = direct.build_prompt({**hybrid_spec(), 'slide_count':0})
    assert 'Exactly 5 pages' in prompt and 'EXACTLY' in prompt and 'claude_content_plan.json' in prompt
    author = claude.build_outline_prompt(value)
    assert 'premise → mechanism → implication' in author and 'full developed Chinese prose' in author
    assert 'Do not create PPTX' in author and 'Never invent statistics' in author


@pytest.mark.parametrize('field', ['audience_cn','central_thesis_cn','narrative_arc_cn','evidence_boundary_cn',
                                 'argument_cn','speaker_notes_cn','transition_cn','source_quotes','evidence_status'])
def test_missing_depth_or_unsupported_evidence_cannot_reach_images(field):
    value = content()
    if field in value: value[field]=''
    else: value['slide_outline'][0][field]=[] if field=='source_quotes' else ''
    with pytest.raises(ValueError): validate_content_plan(value, spec())


def test_codex_can_redesign_layout_but_cannot_erase_the_argument():
    value = content(); slides=copy.deepcopy(value['slide_outline'])
    slides[0]['page_type']='comparison'; slides[0]['layout_description']='new visual explanation'
    validate_content_handoff(slides, value)
    for key in ('visible_text_cn','speaker_notes_cn','source_quotes','evidence_status','action_title_cn'):
        changed=copy.deepcopy(slides); changed[0][key]='substituted'
        with pytest.raises(ValueError):validate_content_handoff(changed, value)


def test_hybrid_packages_notes_and_verified_downloadable_text(tmp_path):
    fixture_output(tmp_path)
    value=hybrid_spec()
    (tmp_path/'deck_plan.json').write_text(json.dumps(content()))
    actual,qa=direct.package_deck('ppt_abcdef0123',tmp_path,value)
    output=tmp_path/'output'
    manifest=verify_published('ppt_abcdef0123','ppt',actual,output)
    assert {'content_outline.md','content_plan.json'} <= {f['path'] for f in manifest['files']}
    assert (output/'content_outline.md').read_text()==outline_markdown(content())
    deck=Presentation(output/'ppt_abcdef0123_deck.pptx')
    assert all(p.notes_slide.notes_text_frame.text==s['speaker_notes_cn'] for p,s in zip(deck.slides,content()['slide_outline']))
    urls=delivery_urls('ppt_abcdef0123',manifest)
    evt={'compiled_spec':actual,'qa_report':qa,'cdn_urls':urls}
    validate_direct_completion(evt,value)
    (output/'content_outline.md').write_text('tampered outline')
    with pytest.raises(ValueError):verify_published('ppt_abcdef0123','ppt',actual,output)


def test_content_handoff_failure_never_publishes(tmp_path):
    fixture_output(tmp_path)
    changed=content();changed['slide_outline'][0]['visible_text_cn']='generic replacement'
    (tmp_path/'deck_plan.json').write_text(json.dumps(changed))
    with pytest.raises(ValueError,match='改写或遗漏'):direct.package_deck('ppt_abcdef0123',tmp_path,hybrid_spec())
    assert not (tmp_path/'output').exists()


@pytest.mark.parametrize('result_kind',['empty','bad_plan','rejected'])
def test_real_planner_failure_never_starts_codex(tmp_path,enabled,monkeypatch,result_kind):
    script=tmp_path/'planner'
    code=''
    if result_kind=='bad_plan':code='from pathlib import Path\nPath("draft/content_plan.json").write_text("{}")\n'
    result={'type':'result','subtype':'success','total_cost_usd':.1,'is_error':False}
    if result_kind=='rejected':result={'type':'system','subtype':'api_retry','error_status':429}
    script.write_text('#!'+sys.executable+'\n'+code+'print('+repr(json.dumps(result))+',flush=True)\n')
    script.chmod(0o700)
    monkeypatch.setattr(claude,'renderer_tools',lambda:pytest.fail('text planning does not need native rendering'))
    async def run():
        req=SimpleNamespace(job_id='ppt_abcdef0123',english_spec=hybrid_spec(),runtime_keys={},max_budget_usd=.85)
        sem=asyncio.Semaphore(1)
        result=[json.loads(x[6:]) async for x in direct.generate_deck_stream(req,tmp_path/'jobs','/usr/bin/false',sem,str(script))]
        assert sem._value==1
        return result
    events=asyncio.run(run())
    assert events[-1]['type']=='error' and not any(e['type']=='completed' for e in events)
    assert not list(tmp_path.rglob('codex_exec.log')) and not list(tmp_path.rglob('delivery.json'))


def test_real_process_handoff_does_not_complete_when_drawing_fails(tmp_path,enabled):
    script=tmp_path/'planner'; value=content(2)
    script.write_text('#!'+sys.executable+'\nfrom pathlib import Path\nPath("draft/content_plan.json").write_text('+repr(json.dumps(value))+')\nprint(\'{"type":"result","subtype":"success","is_error":false,"total_cost_usd":0.12}\')\n')
    script.chmod(0o700)
    async def run():
        req=SimpleNamespace(job_id='ppt_abcdef0123',english_spec=hybrid_spec(2),runtime_keys={},max_budget_usd=.85)
        sem=asyncio.Semaphore(1)
        events=[json.loads(x[6:]) async for x in direct.generate_deck_stream(req,tmp_path/'jobs','/usr/bin/false',sem,str(script))]
        assert sem._value==1
        return events
    events=asyncio.run(run())
    job=tmp_path/'jobs/ppt_abcdef0123'
    assert (job/'content_outline.md').read_text()==outline_markdown(value)
    assert (job/'codex_exec.log').exists() and events[-1]['type']=='error'
    assert any(e.get('stage')=='illustrating' for e in events)
    assert not (job/'output').exists()


def test_hybrid_owner_admission_retry_and_refund_are_one_job(manager,enabled,monkeypatch):
    async def run():
        jobs=PresentationJobs(manager); principal={'user_id':'a','api_key':KEY_A,'scope':'generate'}
        async def ready(_):pass
        calls=[]
        async def fail(job,args):
            calls.append(args);await manager.cancel_refund(job,'drawing failed after outline')
        monkeypatch.setattr(jobs,'require_claude_ready',ready);monkeypatch.setattr(jobs,'execute',fail)
        with pytest.raises(PermissionError):await jobs.create({'user_id':'b','api_key':KEY_B},'test','auto',2,'hybrid-deny-001',execution_mode='claude_codex')
        job=await jobs.create(principal,'完整内容','auto',2,'hybrid-test-001',execution_mode='claude_codex')
        await asyncio.gather(*list(jobs.tasks))
        await jobs.retry(principal,job['job_id'],'hybrid-retry-001');await asyncio.gather(*list(jobs.tasks))
        assert len(calls)==2 and all(x['execution_mode']=='claude_codex' and x['account_owner_user_id']=='a' for x in calls)
        assert balance(manager)==(2000,0)
        with sqlite3.connect(manager.db_path) as db:
            assert db.execute("SELECT count(*) FROM credit_ledger WHERE action='FREEZE'").fetchone()[0]==2
            assert db.execute("SELECT count(*) FROM credit_ledger WHERE action='REFUND'").fetchone()[0]==2
    asyncio.run(run())


def test_hybrid_text_download_is_owner_bound(manager,enabled,monkeypatch):
    async def run():
        job='ppt_abcdef0123';root=manager.jobs_root/job;root.mkdir(parents=True)
        fixture_output(root);(root/'deck_plan.json').write_text(json.dumps(content()))
        actual,qa=direct.package_deck(job,root,hybrid_spec())
        await manager.try_freeze_credits(KEY_A,job,'ppt','原文',60)
        await manager.save_compiled_spec(job,actual)
        urls=delivery_urls(job,json.loads((root/'output/delivery.json').read_text()))
        await manager.confirm_settle(job,.2,urls)
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app),base_url='https://test') as c:
            path=f'/artifacts/{job}/content_outline.md'
            assert (await c.get(path)).status_code==401
            assert (await c.get(path,headers={'Authorization':'Bearer '+KEY_B})).status_code==404
            res=await c.get(path,headers={'Authorization':'Bearer '+KEY_A})
            assert res.status_code==200 and res.text==outline_markdown(content())
            assert 'attachment' in res.headers['content-disposition']
    asyncio.run(run())


@pytest.mark.parametrize('defect',['symlink','hardlink','parent_symlink','oversized'])
def test_unsafe_text_plan_is_rejected(tmp_path,defect):
    from gateway.content_plan import read_plan
    root=tmp_path/'job';root.mkdir();outside=tmp_path/'outside';outside.mkdir()
    candidate=root/'draft/content_plan.json'
    if defect=='parent_symlink':
        (root/'draft').symlink_to(outside,target_is_directory=True);(outside/'content_plan.json').write_text('{}')
    else:
        candidate.parent.mkdir();candidate.write_text('{}')
        if defect=='symlink':
            candidate.unlink();candidate.symlink_to(outside/'content_plan.json');(outside/'content_plan.json').write_text('{}')
        if defect=='hardlink':
            import os
            os.link(candidate,outside/'copy.json')
        if defect=='oversized':candidate.write_bytes(b' '*(2*1024*1024+1))
    with pytest.raises(ValueError):read_plan(candidate,root)


def test_cancelled_outline_stops_cli_and_releases_account(tmp_path,enabled):
    import os
    marker=tmp_path/'planner-pid';script=tmp_path/'planner'
    script.write_text('#!'+sys.executable+'\nimport os,time\nfrom pathlib import Path\nPath('+repr(str(marker))+').write_text(str(os.getpid()))\ntime.sleep(30)\n')
    script.chmod(0o700)
    async def run():
        req=SimpleNamespace(job_id='ppt_abcdef0123',english_spec=hybrid_spec(2),runtime_keys={},max_budget_usd=.85)
        sem=asyncio.Semaphore(1)
        async def consume():
            return [x async for x in direct.generate_deck_stream(req,tmp_path/'jobs','/usr/bin/false',sem,str(script))]
        task=asyncio.create_task(consume())
        for _ in range(200):
            if marker.exists():break
            await asyncio.sleep(.01)
        assert marker.exists()
        task.cancel()
        with pytest.raises(asyncio.CancelledError):await task
        assert sem._value==1
        with pytest.raises(ProcessLookupError):os.kill(int(marker.read_text()),0)
        from worker.claude_account import AccountGate
        async with AccountGate.configured(enabled).lease():pass
        assert not list(tmp_path.rglob('codex_exec.log'))
    asyncio.run(run())


def test_gateway_hybrid_handoff_settles_once_after_both_stages(manager,enabled,monkeypatch,tmp_path):
    from gateway import mcp_server
    author=tmp_path/'author';drawer=tmp_path/'drawer'
    author.write_text('#!'+sys.executable+'\nfrom pathlib import Path\nPath("draft/content_plan.json").write_text('+repr(json.dumps(content()))+')\nprint(\'{"type":"result","subtype":"success","is_error":false,"total_cost_usd":0.12}\')\n')
    drawer.write_text('#!'+sys.executable+'''\nimport json
from pathlib import Path
from PIL import Image
plan=json.loads(Path('claude_content_plan.json').read_text())
Path('deck_plan.json').write_text(json.dumps(plan))
for i in range(1,6):Image.new('RGB',(1672,941),(i*30,20,50)).save(f'staging/assets/slide_preview_{i}.png')
Path('review.json').write_text(json.dumps({'reviewed_pages':[1,2,3,4,5],'issues':[]}))
''')
    author.chmod(0o700);drawer.chmod(0o700)
    class Stream:
        def __init__(self,request):self.request=SimpleNamespace(**request)
        async def __aenter__(self):return self
        async def __aexit__(self,*_):pass
        def raise_for_status(self):pass
        async def aiter_lines(self):
            async for line in direct.generate_deck_stream(self.request,manager.jobs_root,str(drawer),asyncio.Semaphore(1),str(author)):
                yield line.strip()
    class Client:
        def __init__(self,*_,**kw):pass
        async def __aenter__(self):return self
        async def __aexit__(self,*_):pass
        def stream(self,*_,**kw):return Stream(kw['json'])
    monkeypatch.setattr(mcp_server.httpx,'AsyncClient',Client)
    monkeypatch.setattr(mcp_server,'get_runtime_settings',lambda:{'overseas_worker_url':'http://worker','overseas_worker_secret':'test'})
    async def run():
        jobs=PresentationJobs(manager);principal={'user_id':'a','api_key':KEY_A}
        async def ready(_):pass
        monkeypatch.setattr(jobs,'require_claude_ready',ready)
        arguments=dict(source=spec()['original_prompt_cn'],template_id='auto',slide_count=5,request_id='hybrid-settle-01',execution_mode='claude_codex')
        response=await jobs.create(principal,**arguments)
        await asyncio.gather(*list(jobs.tasks))
        status=await jobs.status(principal,response['job_id'])
        assert status['status']=='COMPLETED' and status['credits_held']==0
        assert status['request']['execution_mode']=='claude_codex'
        result=await jobs.artifacts(principal,response['job_id'])
        assert len(result['files'])==9 and result['artifacts']['content_outline_url']
        # Repeating submission reads the completed job and cannot reserve or settle again.
        assert (await jobs.create(principal,**arguments))['status']=='COMPLETED'
        with sqlite3.connect(manager.db_path) as db:
            assert db.execute("SELECT count(*) FROM credit_ledger WHERE action='FREEZE'").fetchone()[0]==1
            assert db.execute("SELECT count(*) FROM credit_ledger WHERE action='SETTLE'").fetchone()[0]==1
            assert db.execute("SELECT count(*) FROM credit_ledger WHERE action='REFUND'").fetchone()[0]==0
    asyncio.run(run())


@pytest.mark.parametrize('mode',['claude_cli','claude_codex'])
def test_worker_enforces_owner_for_every_claude_stage(enabled,monkeypatch,mode):
    from worker import worker_api
    monkeypatch.setattr(worker_api,'WORKER_SECRET','test-worker-'+'x'*40)
    req=worker_api.JobExecutionRequest(job_id='ppt_abcdef0123',task_type='ppt',english_spec={
        'generation_contract':'skill-direct-v1','generator_engine':mode,'account_owner_user_id':'b'},execution_mode=mode)
    async def run():
        with pytest.raises(Exception) as error:await worker_api.execute_stream(req,worker_api.WORKER_SECRET)
        assert error.value.status_code==403
        legacy=worker_api.JobExecutionRequest(job_id='ppt_abcdef0123',task_type='ppt',english_spec={},execution_mode=mode)
        with pytest.raises(Exception) as error:await worker_api.execute_stream(legacy,worker_api.WORKER_SECRET)
        assert error.value.status_code==422
    asyncio.run(run())


@pytest.mark.parametrize('defect',['missing_text','changed_plan','missing_notes'])
def test_hybrid_delivery_cannot_omit_text_or_notes(tmp_path,defect):
    fixture_output(tmp_path);(tmp_path/'deck_plan.json').write_text(json.dumps(content()))
    actual,_=direct.package_deck('ppt_abcdef0123',tmp_path,hybrid_spec());output=tmp_path/'output'
    if defect=='missing_text':(output/'content_outline.md').unlink()
    if defect=='changed_plan':
        changed=content();changed['central_thesis_cn']='different thesis'
        (output/'content_plan.json').write_text(json.dumps(changed))
    if defect=='missing_notes':
        deck=Presentation(output/'ppt_abcdef0123_deck.pptx');deck.slides[0].notes_slide.notes_text_frame.text=''
        deck.save(output/'ppt_abcdef0123_deck.pptx')
    with pytest.raises(ValueError):verify_published('ppt_abcdef0123','ppt',actual,output)


def test_fresh_oauth_does_not_make_a_maintenance_request(tmp_path,enabled,monkeypatch):
    import time
    from worker import claude_account as account
    enabled.credentials.write_text(json.dumps({'claudeAiOauth':{'expiresAt':(time.time()+3600)*1000}}))
    async def forbidden(*_,**kw):pytest.fail('unnecessary maintenance request')
    monkeypatch.setattr(account.asyncio,'create_subprocess_exec',forbidden)
    async def run():
        gate=account.AccountGate.configured(enabled)
        async with gate.lease():assert await account.ensure_fresh_login('/usr/bin/false',enabled,gate)==0
    asyncio.run(run())


@pytest.mark.parametrize('outcome',['success','unchanged','auth_error','rate_limit'])
def test_expired_login_uses_one_toolless_official_probe_and_fails_closed(tmp_path,enabled,outcome):
    import time
    from worker import claude_account as account
    enabled.credentials.write_text(json.dumps({'claudeAiOauth':{'expiresAt':(time.time()-3600)*1000}}))
    script=tmp_path/'official-cli';args=tmp_path/'probe-args.json'
    code='import os,time,json,sys\nfrom pathlib import Path\nPath('+repr(str(args))+').write_text(json.dumps(sys.argv[1:]))\n'
    if outcome=='success':code+='p=Path.home()/".claude/.credentials.json"\nt=p.with_suffix(".tmp")\nt.write_text(json.dumps({"claudeAiOauth":{"expiresAt":(time.time()+3600)*1000}}))\nt.chmod(0o600)\nos.replace(t,p)\n'
    result={'type':'result','subtype':'success','is_error':False,'total_cost_usd':.001}
    if outcome=='auth_error':result.update(is_error=True,result='Failed to refresh OAuth token: private diagnostic not for state')
    if outcome=='rate_limit':result={'type':'system','subtype':'api_retry','error_status':429,'retry_delay_ms':1000}
    script.write_text('#!'+sys.executable+'\n'+code+'print('+repr(json.dumps(result))+',flush=True)\n');script.chmod(0o700)
    async def run():
        gate=account.AccountGate.configured(enabled)
        async with gate.lease():
            if outcome=='success':assert await account.ensure_fresh_login(str(script),enabled,gate)==.001
            else:
                with pytest.raises(account.ProviderUnavailable):await account.ensure_fresh_login(str(script),enabled,gate)
        flags=json.loads(args.read_text())
        assert flags[flags.index('--tools')+1]=='' and '--restricted' in flags
        assert flags[flags.index('--mcp-config')+1]=='{"mcpServers":{}}'
        assert flags[-1]=='Reply only OK. Do not use tools.' and flags[flags.index('--max-budget-usd')+1]=='.05'
        if outcome!='success':
            assert gate.status()['state']==('cooldown' if outcome=='rate_limit' else 'blocked')
            assert 'private diagnostic' not in gate.state_path.read_text()
    asyncio.run(run())
