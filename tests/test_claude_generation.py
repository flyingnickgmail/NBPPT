"""Free creation, owner admission, process control and independent file delivery."""
import asyncio
import hashlib
import json
import os
import sqlite3
import sys
from pathlib import Path
from types import SimpleNamespace
import zipfile

import httpx
import pytest
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from gateway import main
from gateway.deepseek_compiler import compile_cn_request_to_english_spec, validate_direct_completion
from gateway.mcp_server import PresentationJobs
from gateway.native_ppt import validate_native_pptx
from gateway.ppt_contract import ProviderUnavailable
from worker import claude_generation as claude, isolation, native_deck
from worker.claude_account import ClaudeProfile, AccountGate
from test_direct_generation import spec, plan
from test_security_delivery import KEY_A, KEY_B, balance, manager


@pytest.fixture
def enabled(monkeypatch, tmp_path):
    monkeypatch.setenv('NBPPT_CLAUDE_ENABLED', 'true')
    monkeypatch.setenv('NBPPT_CLAUDE_OWNER_USER_ID', 'a')
    monkeypatch.setenv('NBPPT_ACCOUNT_STATE_DIR', str(tmp_path / 'accounts'))
    monkeypatch.setenv('NBPPT_REQUIRE_ISOLATION', 'false')
    home = tmp_path / 'home'
    (home / '.claude').mkdir(parents=True, mode=0o700)
    (home / '.claude.json').write_text(json.dumps({'oauthAccount': {'accountUuid': 'unit-test-account'}}))
    auth = home / '.claude/.credentials.json'
    auth.write_text('{"fake-test-only":"not-a-real-token"}')
    auth.chmod(0o600)
    monkeypatch.setenv('NBPPT_CLAUDE_HOME', str(home))
    async def login(*_): pass
    monkeypatch.setattr(claude, 'verify_login', login)
    return ClaudeProfile.configured()


def make_native(path, count=2):
    path.parent.mkdir(parents=True, exist_ok=True)
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(16), Inches(9)
    for i in range(count):
        page = deck.slides.add_slide(deck.slide_layouts[6])
        shape = page.shapes.add_textbox(Inches(1), Inches(1), Inches(12), Inches(2))
        shape.text = f'第{i+1}页：独立论点'
    deck.save(path)
    return path


def test_free_contract_preserves_source_without_preplanned_layout(enabled, monkeypatch):
    monkeypatch.setattr('gateway.deepseek_compiler.get_runtime_settings', lambda: pytest.fail('no legacy planner'))
    value = asyncio.run(compile_cn_request_to_english_spec(spec()['original_prompt_cn'], 'ppt', {'slide_count': 5, 'execution_mode': 'claude_cli'}))
    assert value['generation_contract'] == 'skill-direct-v1'
    assert value['presentation_format'] == 'native-pptx-v1'
    assert value['matched_template']['id'] == 'auto'
    assert value['original_prompt_cn'] == spec()['original_prompt_cn']
    prompt = claude.build_prompt(value)
    assert 'exactly 5 slides' in prompt and 'own the thesis' in prompt
    assert 'no house template' in prompt and 'source-svg' not in prompt
    assert 'brand.json' in claude.build_prompt({**value, 'template_id': 'customer-brand'})
    args = claude.command('/usr/bin/claude', prompt, .85, profile=enabled)
    settings = json.loads(args[args.index('--settings')+1])
    assert settings['sandbox']['failIfUnavailable'] is True
    assert settings['sandbox']['allowUnsandboxedCommands'] is False
    assert settings['sandbox']['network']['allowedDomains'] == []
    assert '--restricted' in args and '--safe-mode' in args
    assert '--dangerously-skip-permissions' not in args
    assert args[args.index('--mcp-config')+1] == '{"mcpServers":{}}'
    assert 'Agent' not in args[args.index('--tools')+1]


@pytest.mark.parametrize('budget', [0, -1, 6, float('nan'), float('inf'), True, '1'])
def test_invalid_budgets_rejected(budget):
    with pytest.raises(ValueError):
        claude.command('/usr/bin/claude', 'test', budget)


@pytest.mark.parametrize('user,key', [('b', KEY_B), ('admin', 'test-admin-'+'c'*40)])
def test_keys_without_claude_permission_are_denied_including_admin(manager, enabled, user, key):
    async def run():
        jobs = PresentationJobs(manager)
        principal = {'user_id': user, 'api_key': key, 'scope': 'admin' if user == 'admin' else 'generate'}
        with pytest.raises(PermissionError):
            await jobs.create(principal, '产品介绍', 'auto', 2, 'other-owner-01', execution_mode='claude_cli')
        assert not jobs.tasks and balance(manager) == (2000, 0)
    asyncio.run(run())


def test_missing_owner_binding_disables_provider_before_freeze(manager, enabled, monkeypatch):
    monkeypatch.delenv('NBPPT_CLAUDE_OWNER_USER_ID')
    async def run():
        jobs = PresentationJobs(manager)
        with pytest.raises(ValueError):
            await jobs.create({'user_id': 'a', 'api_key': KEY_A, 'scope': 'generate'}, 'test', 'auto', 2, 'no-owner-001', execution_mode='claude_cli')
        assert balance(manager) == (2000, 0)
    asyncio.run(run())


def test_circuit_rejected_before_freeze_and_rest_returns_retry_after(manager, enabled, monkeypatch):
    async def run():
        jobs = PresentationJobs(manager)
        async def paused(_): return {'state': 'cooldown', 'message': '冷却中', 'retry_after_seconds': 60}
        monkeypatch.setattr(jobs, 'claude_status', paused)
        monkeypatch.setattr(main, 'mcp_jobs', jobs)
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app), base_url='https://test') as client:
            res = await client.post('/api/jobs', headers={'Authorization': 'Bearer '+KEY_A}, json={'source':'产品介绍','slide_count':2,'execution_mode':'claude_cli','request_id':'cooldown-test-01'})
            assert res.status_code == 429 and res.headers['retry-after'] == '60'
            assert res.json()['detail']['provider_state'] == 'cooldown'
            legacy = await client.post('/api/jobs/stream', headers={'Authorization': 'Bearer '+KEY_A}, json={'prompt_cn':'test','slide_count':2,'execution_mode':'claude_cli'})
            assert legacy.status_code == 429
        assert balance(manager) == (2000, 0) and not jobs.tasks
        with sqlite3.connect(manager.db_path) as db:
            assert db.execute('SELECT count(*) FROM credit_ledger').fetchone()[0] == 0
    asyncio.run(run())


def test_claude_retry_preserves_provider_owner_source_and_refunds(manager, enabled, monkeypatch):
    async def run():
        jobs = PresentationJobs(manager)
        calls = []
        async def ready(_): return None
        async def fail(job, args):
            calls.append(args)
            await manager.cancel_refund(job, 'controlled Claude failure')
        monkeypatch.setattr(jobs, 'require_claude_ready', ready)
        monkeypatch.setattr(jobs, 'execute', fail)
        principal = {'user_id':'a','api_key':KEY_A,'scope':'generate'}
        created = await jobs.create(principal, '产品原文', 'auto', 2, 'claude-create-01', execution_mode='claude_cli')
        await asyncio.gather(*list(jobs.tasks))
        with pytest.raises(ValueError):
            await jobs.create(principal, '产品原文', 'auto', 2, 'claude-create-01', execution_mode='codex_cli')
        await jobs.retry(principal, created['job_id'], 'claude-retry-01')
        await asyncio.gather(*list(jobs.tasks))
        assert len(calls) == 2 and all(x['execution_mode']=='claude_cli' and x['account_owner_user_id']=='a' and x['source']=='产品原文' for x in calls)
        assert balance(manager) == (2000, 0)
        with sqlite3.connect(manager.db_path) as db:
            assert db.execute("SELECT count(*) FROM credit_ledger WHERE action='REFUND'").fetchone()[0] == 2
    asyncio.run(run())


def test_video_cannot_silently_substitute_claude_with_codex(manager, enabled):
    async def run():
        with pytest.raises(ValueError, match='视频尚未开放 Claude'):
            await PresentationJobs(manager).create({'user_id':'a','api_key':KEY_A,'scope':'generate'}, 'test','auto',0,'claude-video-01',task_type='video',execution_mode='claude_cli')
        assert balance(manager) == (2000, 0)
    asyncio.run(run())


def test_linux_bind_uses_one_canonical_directory_not_job_token_copies(tmp_path, enabled, monkeypatch):
    monkeypatch.setenv('NBPPT_REQUIRE_ISOLATION','true')
    monkeypatch.setattr(isolation.sys,'platform','linux')
    monkeypatch.setattr(isolation.os,'geteuid',lambda:1001)
    monkeypatch.setattr(isolation.shutil,'which',lambda _: '/usr/bin/bwrap')
    job=tmp_path/'job'; job.mkdir()
    args, sandbox = isolation.isolated_command(['/usr/bin/claude'],job,[],agent='claude',profile=enabled)
    assert str(enabled.credentials.parent) in args and str(enabled.credentials) not in args
    assert not (sandbox/'.claude/.credentials.json').exists()
    assert '--clearenv' in args and 'CLAUDE_CODE_MAX_RETRIES' in args
    render, _ = isolation.isolated_command(['/usr/bin/soffice'],job,[],agent='renderer')
    assert '--unshare-net' in render and str(enabled.credentials.parent) not in render


@pytest.mark.parametrize('encoding', ['utf-8', 'utf-16', 'utf-16-le', 'utf-16-be', 'utf-32-le', 'utf-32-be'])
def test_native_gate_rejects_entities_in_every_xml_encoding(tmp_path, encoding):
    p = tmp_path / 'entity.pptx'
    make_native(p)
    with zipfile.ZipFile(p, 'a') as z:
        z.writestr('test.xml', '<!DOCTYPE a [<!ENTITY x "expanded">]><a>&x;</a>'.encode(encoding))
    with pytest.raises(ValueError, match='文档实体'):
        validate_native_pptx(p)


@pytest.mark.parametrize('defect', ['corrupt','count','size','blank','external','entity','macro','traversal','symlink','hardlink'])
def test_bad_native_deck_cannot_pass_file_gate(tmp_path, defect):
    p=make_native(tmp_path/'draft/presentation.pptx')
    if defect=='corrupt': p.write_bytes(b'not a PPTX')
    if defect=='count':
        with pytest.raises(ValueError): validate_native_pptx(p,3)
        return
    if defect in ('size','blank'):
        deck=Presentation(p)
        if defect=='size': deck.slide_width=deck.slide_height
        else:
            deck.slides.add_slide(deck.slide_layouts[6])
        deck.save(p)
    if defect in ('external','entity','macro','traversal'):
        with zipfile.ZipFile(p,'a') as z:
            if defect=='external': z.writestr('test.rels','<Relationships><Relationship TargetMode="External" Target="file:///private/secret"/></Relationships>')
            if defect=='entity': z.writestr('test.xml','<!DOCTYPE a [<!ENTITY x "evil">]><a/>')
            if defect=='macro': z.writestr('ppt/vbaProject.bin',b'macro')
            if defect=='traversal': z.writestr('../outside',b'bad')
    if defect in ('symlink','hardlink'):
        other=tmp_path/'outside.pptx'; p.rename(other)
        if defect=='symlink': p.symlink_to(other)
        else: os.link(other,p)
    with pytest.raises((ValueError, zipfile.BadZipFile)): validate_native_pptx(p)


def test_native_pptx_is_preserved_and_previews_are_rendered(tmp_path, monkeypatch):
    job=tmp_path/'job'; p=make_native(job/'draft/presentation.pptx')
    original=hashlib.sha256(p.read_bytes()).hexdigest()
    monkeypatch.setattr(native_deck,'renderer_tools',lambda:{'soffice':'soffice','pdftoppm':'pdftoppm'})
    async def renderer(args, root):
        render=root/'rendered'
        if args[0]=='soffice': (render/'presentation.pdf').write_bytes(b'unit-test pdf')
        else:
            for i in (1,2):
                im=Image.new('RGB',(1672,941),'white');ImageDraw.Draw(im).rectangle((100,100,200*i,400),fill='navy')
                im.save(render/f'page-{i}.png')
    monkeypatch.setattr(native_deck,'run_renderer',renderer)
    actual,qa=asyncio.run(native_deck.package_native('ppt_c1a0de0001',job,{'slide_count':2}))
    assert actual['slide_count']==2 and actual['slide_outline'][0]['evidence_status']=='NOT_REVIEWED'
    assert qa['status']=='PASSED'
    assert hashlib.sha256((job/'output/ppt_c1a0de0001_deck.pptx').read_bytes()).hexdigest()==original
    manifest=json.loads((job/'output/delivery.json').read_text())
    assert manifest['content_status']=='NOT_REVIEWED' and len(list((job/'output/assets').glob('*.png')))==2


@pytest.mark.parametrize('result', [None, {'type':'result','subtype':'error_max_budget_usd','is_error':True}, {'type':'result','subtype':'success','is_error':False}])
def test_real_process_without_success_report_cannot_complete(tmp_path, monkeypatch, enabled, result):
    script=tmp_path/'fake-cli'
    script.write_text('#!'+sys.executable+'\nimport json\nprint('+repr(json.dumps(result) if result else 'nothing')+')\n')
    script.chmod(0o700)
    monkeypatch.setattr(claude,'renderer_tools',lambda: {})
    async def run():
        req=SimpleNamespace(job_id='ppt_c1a0de0002',english_spec={**spec(),'account_owner_user_id':'a','generator_engine':'claude_cli'},max_budget_usd=.85)
        sem=asyncio.Semaphore(1)
        events=[json.loads(e[6:]) async for e in claude.generate_deck_stream(req,tmp_path/'jobs',str(script),sem)]
        assert sem._value==1
        return events
    events=asyncio.run(run())
    assert events[-1]['type']=='error' and not any(x['type']=='completed' for x in events)
    assert (tmp_path/'jobs/ppt_c1a0de0002/claude_exec.log').exists()  # reached the actual process
    assert not list((tmp_path/'jobs').rglob('delivery.json'))


def test_rate_limit_kills_cli_before_retry_and_persists_circuit(tmp_path, monkeypatch, enabled):
    script=tmp_path/'fake-cli'; marker=tmp_path/'should-never-run'
    script.write_text('#!'+sys.executable+'\nimport time,json,pathlib,signal\nsignal.signal(signal.SIGTERM,signal.SIG_IGN)\nprint(json.dumps({"type":"system","subtype":"api_retry","error_status":429,"retry_delay_ms":1000}),flush=True)\ntime.sleep(1)\npathlib.Path('+repr(str(marker))+').write_text("retried")\n')
    script.chmod(0o700)
    monkeypatch.setattr(claude,'renderer_tools',lambda: {})
    async def run():
        req=SimpleNamespace(job_id='ppt_c1a0de0003',english_spec={**spec(),'account_owner_user_id':'a'},max_budget_usd=.85)
        return [json.loads(e[6:]) async for e in claude.generate_deck_stream(req,tmp_path/'jobs',str(script),asyncio.Semaphore(1))]
    events=asyncio.run(run())
    assert events[-1]['provider_state']=='cooldown'
    assert not marker.exists()
    assert AccountGate(enabled.account_id).status()['state']=='cooldown'
    assert not list((tmp_path/'jobs').rglob('delivery.json'))


def test_completion_cannot_change_admitted_provider():
    expected={**spec(),'generation_contract':'skill-direct-v1','generator_engine':'claude_cli','presentation_format':'native-pptx-v1'}
    event={'compiled_spec':{**expected,**plan(),'generator_engine':'codex_cli'},'qa_report':{'status':'PASSED','planned_slides':5,'delivered_slides':5,'pptx_slides_count':5},'cdn_urls':{'pptx_url':'/a.pptx','html_deck_url':'/a.html','slide_preview_urls':[f'/{i}.png' for i in range(5)]}}
    with pytest.raises((ValueError,RuntimeError)): validate_direct_completion(event,expected)

@pytest.mark.parametrize('bad_render',['missing','blank','duplicate','wrong_count'])
def test_render_failure_cannot_publish_even_with_valid_pptx(tmp_path,monkeypatch,bad_render):
    job=tmp_path/'job';make_native(job/'draft/presentation.pptx')
    monkeypatch.setattr(native_deck,'renderer_tools',lambda:{'soffice':'soffice','pdftoppm':'pdftoppm'})
    async def renderer(args,root):
        render=root/'rendered'
        if args[0]=='soffice':
            if bad_render!='missing':(render/'presentation.pdf').write_bytes(b'test')
        else:
            for i in range(1,2 if bad_render=='wrong_count' else 3):
                im=Image.new('RGB',(1672,941),'white')
                if bad_render!='blank':ImageDraw.Draw(im).rectangle((10,10,100 if bad_render=='duplicate' else 100*i,300),fill='black')
                im.save(render/f'page-{i}.png')
    monkeypatch.setattr(native_deck,'run_renderer',renderer)
    with pytest.raises(ValueError):asyncio.run(native_deck.package_native('ppt_c1a0de0004',job,{'slide_count':2}))
    assert not (job/'output').exists()


def test_legacy_worker_cannot_bypass_owner_gate(enabled,monkeypatch):
    from worker import worker_api
    monkeypatch.setattr(worker_api,'WORKER_SECRET','test-worker-'+'x'*40)
    req=worker_api.JobExecutionRequest(job_id='ppt_c1a0de0005',task_type='ppt',english_spec={},execution_mode='claude_cli')
    async def run():
        with pytest.raises(Exception) as caught:await worker_api.execute_stream(req,worker_api.WORKER_SECRET)
        assert caught.value.status_code==422
    asyncio.run(run())


def test_worker_rate_limit_refunds_once_even_if_failure_annotation_fails(manager,enabled,monkeypatch):
    from gateway import mcp_server
    async def run():
        jobs=PresentationJobs(manager);job='ppt_c1a0de0006'
        await manager.try_freeze_credits(KEY_A,job,'ppt','原文',60)
        class Stream:
            async def __aenter__(self):return self
            async def __aexit__(self,*_):pass
            def raise_for_status(self):pass
            async def aiter_lines(self):yield 'data: {"type":"error","provider_state":"cooldown","retry_after_seconds":60}'
        class Client:
            def __init__(self,*_,**kw):pass
            async def __aenter__(self):return self
            async def __aexit__(self,*_):pass
            def stream(self,*_,**kw):return Stream()
        monkeypatch.setattr(mcp_server.httpx,'AsyncClient',Client)
        original_get=manager.get_job
        async def fail_annotation(*_):raise RuntimeError('annotation lookup failed')
        monkeypatch.setattr(manager,'get_job',fail_annotation)
        with pytest.raises(RuntimeError,match='annotation lookup'):
            await jobs.execute(job,{'source':'原文','template_id':'auto','slide_count':2,'execution_mode':'claude_cli','account_owner_user_id':'a'})
        monkeypatch.setattr(manager,'get_job',original_get)
        assert balance(manager)==(2000,0)
        await manager.cancel_refund(job,'duplicate cleanup')
        with sqlite3.connect(manager.db_path) as db:
            assert db.execute("SELECT count(*) FROM credit_ledger WHERE action='REFUND'").fetchone()[0]==1
            assert db.execute('SELECT status FROM media_jobs').fetchone()[0]=='FAILED'
    asyncio.run(run())

@pytest.mark.parametrize('name,content', [
 ('test.rels','<Relationships><Relationship Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="harmless.bin"/></Relationships>'),
 ('ppt/hidden.xml','<root ContentType="application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml"/>'),
 ('ppt/media/image.svg','<svg xmlns="http://www.w3.org/2000/svg"><image href="file:///private/secret"/></svg>'),
 ('ppt/media/image.svg','<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>')])
def test_active_content_cannot_hide_behind_an_innocent_filename(tmp_path,name,content):
    p=make_native(tmp_path/'draft/presentation.pptx')
    with zipfile.ZipFile(p,'a') as z:z.writestr(name,content)
    with pytest.raises(ValueError):validate_native_pptx(p)


def test_rejects_job_storage_inside_login_home(tmp_path,enabled):
    with pytest.raises(ValueError):native_deck.validate_job_location(enabled.home/'jobs',enabled)


def test_capabilities_require_key_grant_and_worker_requires_runtime_binding(manager,enabled,monkeypatch):
    from worker import worker_api
    async def run():
        jobs=PresentationJobs(manager)
        async def ready(_):return {'state':'ready','retry_after_seconds':0,'message':'ready'}
        monkeypatch.setattr(jobs,'claude_status',ready);monkeypatch.setattr(main,'mcp_jobs',jobs)
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app),base_url='https://test') as c:
            other=await c.get('/api/capabilities',headers={'Authorization':'Bearer '+KEY_B})
            assert other.json()['delivery_modes']['ppt']['execution_modes']==['codex_cli']
            denied=await c.post('/api/jobs',headers={'Authorization':'Bearer '+KEY_B},json={'source':'未授权的密钥不能使用Claude','slide_count':2,'execution_mode':'claude_cli','request_id':'owner-bypass-001'})
            assert denied.status_code==403
        monkeypatch.setattr(worker_api,'WORKER_SECRET','test-worker-'+'x'*40)
        req=worker_api.JobExecutionRequest(job_id='ppt_c1a0de0008',task_type='ppt',english_spec={'generation_contract':'skill-direct-v1','generator_engine':'claude_cli','account_owner_user_id':'b'},execution_mode='claude_cli')
        with pytest.raises(Exception) as caught:await worker_api.execute_stream(req,worker_api.WORKER_SECRET)
        assert caught.value.status_code==403 and balance(manager)==(2000,0)
    asyncio.run(run())


def test_rate_limit_stops_orphaned_shell_children_after_cli_exit(tmp_path,enabled,monkeypatch):
    marker=tmp_path/'orphan-retry';script=tmp_path/'fake-cli'
    child='import time,pathlib;time.sleep(1);pathlib.Path('+repr(str(marker))+').write_text("unsafe retry")'
    script.write_text('#!'+sys.executable+'\nimport sys,subprocess,json\nsubprocess.Popen([sys.executable,"-c",'+repr(child)+'],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)\nprint(json.dumps({"type":"system","subtype":"api_retry","error_status":429}),flush=True)\n')
    script.chmod(0o700);monkeypatch.setattr(claude,'renderer_tools',lambda: {})
    async def run():
        req=SimpleNamespace(job_id='ppt_c1a0de0009',english_spec={**spec(),'account_owner_user_id':'a'},max_budget_usd=.85)
        events=[json.loads(e[6:]) async for e in claude.generate_deck_stream(req,tmp_path/'jobs',str(script),asyncio.Semaphore(1))]
        assert events[-1]['provider_state']=='cooldown'
        await asyncio.sleep(1.2)
        assert not marker.exists()
    asyncio.run(run())


def test_oversized_xml_rejected_before_parsing(tmp_path):
    p=make_native(tmp_path/'draft/presentation.pptx')
    with zipfile.ZipFile(p,'a',compression=zipfile.ZIP_DEFLATED) as z:z.writestr('large.xml','<a>'+('x'* (4*1024*1024))+'</a>')
    with pytest.raises(ValueError,match='XML'):validate_native_pptx(p)


@pytest.mark.parametrize('family',['Heiti SC','Noto Sans CJK SC'])
def test_missing_required_cjk_font_is_not_silently_substituted(monkeypatch,family):
    monkeypatch.setenv('NBPPT_SOFFICE',sys.executable);monkeypatch.setenv('NBPPT_PDFTOPPM',sys.executable)
    monkeypatch.setattr(native_deck.shutil,'which',lambda _:sys.executable)
    monkeypatch.setattr(native_deck.subprocess,'run',lambda *a,**kw:SimpleNamespace(returncode=0,stdout=family.encode()))
    if family.startswith('Noto'):assert 'soffice' in native_deck.renderer_tools()
    else:
        with pytest.raises(ValueError,match='Noto'):native_deck.renderer_tools()

@pytest.mark.parametrize('body',['<style>@import url(https://example.test/x.css)</style>','<g xml:base="file:///private/"/>'])
def test_embedded_svg_cannot_hide_external_loading_in_styles_or_base(tmp_path,body):
    p=make_native(tmp_path/'draft/presentation.pptx')
    with zipfile.ZipFile(p,'a') as z:z.writestr('ppt/media/image.svg','<svg xmlns="http://www.w3.org/2000/svg">'+body+'</svg>')
    with pytest.raises(ValueError):validate_native_pptx(p)
